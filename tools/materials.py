"""Materials list + cut list workbook, written from takeoff.py.

The takeoff itself — every solid that gets cut, the board packing, the sheet nest —
lives in takeoff.py at the repo root, because the v2 site draws the same nest on its
Cut list page and a second copy would drift. This module is only the workbook: tab
layout, cell styling, and the TBD/assumptions narrative, which is a purchasing
document rather than a property of the model.

Development harness, like everything else in tools/: nothing here is imported by the
build, and the workbook is a deliverable, not a source of truth. Run it again after
any change to the model.

    python3 tools/materials.py            -> materials/materials-list.xlsx
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from takeoff import layout_svg, takeoff

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "materials", "materials-list.xlsx")

THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="FF1F3864")
BAND = PatternFill("solid", fgColor="FFF2F2F2")
FLAG = PatternFill("solid", fgColor="FFFFF2CC")

# ------------------------------------------------------------------ workbook
def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFFFF", size=11)
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = f"A{row + 1}"


def widths(ws, spec):
    for i, w in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write(ws, rows, header, wide):
    ws.append(header)
    style_header(ws, len(header))
    for i, r in enumerate(rows):
        ws.append(r)
        if i % 2:
            for c in range(1, len(header) + 1):
                ws.cell(row=ws.max_row, column=c).fill = BAND
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(header)):
        for cell in row:
            cell.border = BORDER
            if cell.row > 1:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths(ws, wide)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"


def main():
    t = takeoff()
    wb = Workbook()

    ws = wb.active
    ws.title = "Buy list"
    write(ws, t["buy"],
          ["Stock", "Species / grade", "Length", "Qty", "Unit", "Notes"],
          [20, 22, 18, 8, 10, 72])

    ws = wb.create_sheet("Cut list")
    write(ws, t["cut"],
          ["Part", "Qty", "Assembly", "Stock", "Length", "Width", "Thick", "Source", "Notes"],
          [26, 6, 12, 18, 10, 10, 8, 30, 78])

    ws = wb.create_sheet("Sheet goods")
    write(ws, t["sheet_goods"],
          ["Panel", "Stock", "Long", "Short", "From a split?", "Notes / joint"],
          [26, 14, 10, 10, 14, 90])

    ws = wb.create_sheet("Cut layout")
    write(ws, t["layout"],
          ["Grade", "Sheet", "Piece", "W", "H", "From edge X", "From edge Y",
           "Grain", "Flag"], [22, 7, 26, 9, 9, 12, 12, 10, 14])
    layout_svg(t["nests"], os.path.join(os.path.dirname(OUT), "cut-layout.svg"))

    ws = wb.create_sheet("Hardware")
    write(ws, t["hardware"], ["Connection / item", "Qty", "Type", "Fastener spec"],
          [46, 8, 14, 110])

    silent = t["silent"]
    # ---------------------------------------------------------- TBD
    ws = wb.create_sheet("TBD and assumptions")
    tbd = [
        ["CLOSED", "Deck and ceiling ply layouts, chosen for joinery (Rev AX)",
         "The deck is THREE pieces: two rows to x 62 1/4 either side of a blocked "
         "cross-joist seam at y 24, then one full-depth 44 x 48 1/4 panel over the landing "
         "end with no seam in it — the part walked on rather than slept on. Both seams sit "
         "under the mattress; the long one lands on joist 5 and needs no blocking. 62 1/4 "
         "is the furthest in it can go: the next joist would need a 56 wide panel against a "
         "48 sheet width. The ceiling is three pieces with BOTH seams on joist centres "
         "(38 1/4, 74 1/4), so every ceiling seam is backed. "
         "Trimming the deck to a flat 48 was considered and rejected — it changed no sheet "
         "count and would have left beam_wrap_inner's foot over a 1/4 void. "
         "COST, stated plainly: these layouts are 9 sheets at 64%. The cheapest nest was 7 "
         "at 83%, with a full-width deck seam and a cross-joist seam in the ceiling. Two "
         "sheets bought better joinery, knowingly."],
        ["CLOSED", "The loft deck IS the finished walking surface (2026-09-07)",
         "The mattress covers y 8 3/4 to 46 3/4 out to about x 77; the rest is walked on "
         "and nothing in the model covers it, so the deck is the finished floor and takes a "
         "sanded paintable face rather than rated sheathing. Consequence: there is no hidden "
         "3/4 ply left in the build, so the two grade groups collapsed to ONE paint-grade "
         "pool — and the sheet count assumes that. Buying the deck as sheathing again would "
         "break the nest, not just change the price."],
        ["GRADE", "Every species/grade cell says TBD",
         "Pending price and availability. TWO ply grades, not three: one 3/4 paint-grade "
         "pool covering every 3/4 piece including the deck, and 1/2 flitches for the jamb "
         "packs and the header's middle ply. Framing species/grade and the slat stock are "
         "also open."],
        ["CLOSED", "Deck-ply front-edge blocking is now in the model (Rev AU)",
         "Nine 2x4 blocks on edge at the beam face — eight at 10 1/2 on a 12 pitch plus "
         "one 5 3/4 at the rim, all from one 8 ft board, and they are counted in the Buy "
         "list and the Cut list. Fixed with 2 x 1/4 x 3 per block driven from the beam's "
         "OUTBOARD face before beam_wrap_face closes it, then the ply glued and screwed "
         "down onto them at 6 in o.c. SEQUENCE: before the deck ply."],
        ["CLOSED", "The beam notch needs no separate ply cheek (Rev AT)",
         "It is beam_wrap_face, which already spans the corner on the clear outboard "
         "face: adhesive + #8 x 2 at 6 in o.c. over x 0-27, into the tongue above z 57 1/4 "
         "and into the ledger and joist-tail end faces below it. Nothing extra to buy — but "
         "the wrap's fixing over that first 27 in is STRUCTURAL, not finish, and its screws "
         "must stay 1 in clear of the corner and out of the notch void below z 57 1/4."],
        ["FIELD", "Stud locations: bed wall, window wall, right wall",
         "Every ledger fixing count on the Hardware tab assumes 16 in o.c. The window "
         "wall needs a stud (or added blocking) within ~6 in of y 50 for the beam's end "
         "reaction."],
        ["FIELD", "Ceiling joist nearest 50 3/4 from the bed wall",
         "No longer gates the design (Rev AS: the plate is fixed at 47 1/4-50 3/4 and "
         "blocking can be added from the attic), but it decides whether you screw to a "
         "joist or block first."],
        ["FIELD", "Floor joists under the half-wall plate and the kicker", ""],
        ["SPECIFY", "hw_bottom_plate_a / _b to floor", "Left TBD by the user 2026-09-06."],
        ["CHECKER", f"{len(silent)} connections carry NO fastener key at all",
         "Filter the Hardware tab for 'NO FASTENER KEY'. verify.py only reports a row "
         "that literally says UNSPECIFIED, so these pass silently: "
         + "; ".join(silent) + ". Most are ordinary nailing, but the kicker is already an "
         "open anchorage item. Worth making verify.py WARN on a missing key — it would "
         "change the 0 WARN baseline, so it is your call. (Counted from the model, not "
         "typed: an earlier hand count said 8.)"],
        ["SPECIFY", "kicker anchorage", ""],
        ["METHOD", "Sheet counts are a real nest, with kerf",
         "Every piece is placed by a guillotine first-fit-decreasing nest that reserves "
         "1/8 on every cut, so the Cut layout tab and materials/cut-layout.svg are a "
         "layout you can mark out, not an indicative count — the diagrams draw each piece "
         "at its finished size with the kerf as the gap between them, and no two pieces "
         "overlap or run off a sheet (asserted at build time). Rotation is allowed and the "
         "layout flags which pieces came out turned, so an appearance call can override "
         "one. TWO CAVEATS. There is no factory-edge trim allowance: the first piece sits "
         "in the corner, so the layout assumes two usable edges — kerf IS reserved against "
         "the far edge, which is roughly the 1/8 back. And the deck and the loft ceiling "
         "are OVER 48 in y (48 1/4 and 50), so that dimension has to run along the sheet's "
         "8 ft length and each takes one sheet per piece."],
        ["METHOD", "Cut lengths are FRAMING extents",
         "Finished faces are their own parts. No waste allowance is added anywhere; "
         "the board counts include a 1/8 kerf per cut and nothing else."],
    ]
    tbd.insert(0, ["SOURCE", f"Generated from dimensions.yaml rev {t['rev']}",
                   "tools/materials.py, off the same model verify.py and the CAD kernel "
                   "read. Regenerate after ANY model change — nothing here updates itself, "
                   "and a stale workbook looks exactly like a current one."])
    write(ws, tbd, ["Kind", "Item", "Detail"], [12, 52, 100])
    wb.properties.title = f"Loft bed materials and cut list — rev {t['rev']}"
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=1):
        row[0].fill = FLAG

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"{OUT}: {t['pieces']} pieces, {len(t['buy'])} buy lines, rev {t['rev']}")


if __name__ == "__main__":
    main()
